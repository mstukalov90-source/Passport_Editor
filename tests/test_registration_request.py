"""Тесты публичной формы заявки на регистрацию и списка заявок для МГГТ."""

from __future__ import annotations

from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook
from pass_viewer import registration_reference, registration_views
from pass_viewer.models import ExternalUser, RegistrationRequest

VALID_PAYLOAD = {
    'executive_authority': 'Минстрой России',
    'institution_name': 'ГУП «Тестовое учреждение»',
    'representative_name': 'Иванов Иван Иванович',
    'position': 'Главный специалист отдела информатизации',
    'phone': '+7 (495) 123-45-67',
    'email': 'ivanov@example.ru',
}


def _login_as(client, login, role):
    ExternalUser.objects.create(login=login, password='pass', role=role)
    client.post(reverse('login'), {'username': login, 'password': 'pass'})


@pytest.fixture(autouse=True)
def _no_reference_lookups(monkeypatch):
    """Тесты страницы заявки не ходят в удалённую БД mggt_asu."""
    monkeypatch.setattr(registration_views, 'list_executive_authorities', lambda: [])
    monkeypatch.setattr(registration_views, 'list_institutions', lambda: [])


@pytest.mark.django_db
def test_registration_form_page_shows_template_fields(client):
    response = client.get(reverse('registration_request'))
    assert response.status_code == 200
    content = response.content.decode('utf-8')
    for label in (
        'Орган исполнительной власти',
        'Наименование учреждения',
        'ФИО ответственного представителя',
        'Должность',
        'Контактный телефон',
        'Адрес электронной почты',
    ):
        assert label in content
    assert 'Отправить заявку' in content


@pytest.mark.django_db
def test_registration_submit_creates_request(client):
    response = client.post(reverse('registration_request'), VALID_PAYLOAD)
    assert response.status_code == 302
    assert response.url == reverse('registration_request_sent')
    assert RegistrationRequest.objects.count() == 1
    item = RegistrationRequest.objects.get()
    assert item.executive_authority == 'Минстрой России'
    assert item.institution_name == 'ГУП «Тестовое учреждение»'
    assert item.representative_name == 'Иванов Иван Иванович'
    assert item.position == 'Главный специалист отдела информатизации'
    assert item.phone == '+7 (495) 123-45-67'
    assert item.email == 'ivanov@example.ru'
    assert item.status == RegistrationRequest.STATUS_NEW
    assert client.get(reverse('registration_request_sent')).status_code == 200


@pytest.mark.django_db
def test_registration_submit_strips_values(client):
    payload = {key: f'  {value}  ' for key, value in VALID_PAYLOAD.items()}
    client.post(reverse('registration_request'), payload)
    item = RegistrationRequest.objects.get()
    assert item.representative_name == 'Иванов Иван Иванович'
    assert item.email == 'ivanov@example.ru'


@pytest.mark.django_db
def test_registration_submit_invalid_shows_russian_errors(client):
    payload = {**VALID_PAYLOAD, 'position': '', 'email': 'not-an-email'}
    response = client.post(reverse('registration_request'), payload)
    assert response.status_code == 200
    assert RegistrationRequest.objects.count() == 0
    content = response.content.decode('utf-8')
    assert 'Заполните это поле.' in content
    assert 'Укажите корректный адрес электронной почты.' in content


@pytest.mark.django_db
def test_login_page_links_to_registration_form(client):
    content = client.get(reverse('login')).content.decode('utf-8')
    assert 'Регистрация пользователя' in content
    assert reverse('registration_request') in content


@pytest.mark.django_db
def test_requests_list_anonymous_redirects_to_login(client):
    response = client.get(reverse('registration_requests_list'))
    assert response.status_code == 302
    assert '/accounts/login/' in response.url


@pytest.mark.django_db
def test_requests_list_forbidden_for_non_mggt(client):
    _login_as(client, 'bd_user', role=ExternalUser.ROLE_BD)
    assert client.get(reverse('registration_requests_list')).status_code == 403


@pytest.mark.django_db
def test_requests_list_shows_items_and_filters_for_mggt(client):
    _login_as(client, 'mggt_list', role=ExternalUser.ROLE_MGGT)
    new_item = RegistrationRequest.objects.create(**VALID_PAYLOAD)
    RegistrationRequest.objects.create(
        **{**VALID_PAYLOAD, 'email': 'processed@example.ru'},
        status=RegistrationRequest.STATUS_PROCESSED,
    )

    response = client.get(reverse('registration_requests_list'))
    assert response.status_code == 200
    content = response.content.decode('utf-8')
    assert 'Заявки на регистрацию' in content
    assert new_item.representative_name in content
    assert 'Новые (1)' in content

    response = client.get(reverse('registration_requests_list'), {'status': 'new'})
    content = response.content.decode('utf-8')
    assert 'ivanov@example.ru' in content
    assert 'processed@example.ru' not in content


@pytest.mark.django_db
def test_set_status_toggles_request_state(client):
    _login_as(client, 'mggt_status', role=ExternalUser.ROLE_MGGT)
    item = RegistrationRequest.objects.create(**VALID_PAYLOAD)
    response = client.post(
        reverse('registration_request_set_status', args=[item.pk]),
        {'status': RegistrationRequest.STATUS_PROCESSED},
    )
    assert response.status_code == 302
    item.refresh_from_db()
    assert item.status == RegistrationRequest.STATUS_PROCESSED


@pytest.mark.django_db
def test_export_xlsx_matches_template(client):
    _login_as(client, 'mggt_export', role=ExternalUser.ROLE_MGGT)
    RegistrationRequest.objects.create(**VALID_PAYLOAD)

    response = client.get(reverse('registration_requests_export'))
    assert response.status_code == 200
    assert response['Content-Type'] == (
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert sheet.title == 'Перечень'
    headers = [cell.value for cell in sheet[1]]
    assert headers == [
        '№ п/п',
        'Орган исполнительной власти',
        'Наименование учреждения',
        'ФИО ответственного представителя',
        'Должность',
        'Контактный телефон',
        'Адрес электронной почты',
    ]
    assert sheet.max_row == 2
    first_row = [cell.value for cell in sheet[2]]
    assert first_row == [
        1,
        'Минстрой России',
        'ГУП «Тестовое учреждение»',
        'Иванов Иван Иванович',
        'Главный специалист отдела информатизации',
        '+7 (495) 123-45-67',
        'ivanov@example.ru',
    ]


@pytest.mark.django_db
def test_export_forbidden_for_non_mggt(client):
    _login_as(client, 'bd_export', role=ExternalUser.ROLE_BD)
    assert client.get(reverse('registration_requests_export')).status_code == 403


@pytest.mark.django_db
def test_form_page_renders_datalist_options(client, monkeypatch):
    monkeypatch.setattr(
        registration_views,
        'list_executive_authorities',
        lambda: [
            {'shortname': 'ДФМ', 'fullname': 'Департамент финансов города Москвы'},
            {'shortname': 'ДЖКХ', 'fullname': 'Департамент жилищно-коммунального хозяйства'},
        ],
    )
    monkeypatch.setattr(
        registration_views,
        'list_institutions',
        lambda: [{'shortname': 'Жилищник', 'fullname': 'ГБУ «Жилищник»'}],
    )
    response = client.get(reverse('registration_request'))
    assert response.status_code == 200
    content = response.content.decode('utf-8')
    assert 'list="registration-executive-authority-options"' in content
    assert 'list="registration-institution-options"' in content
    assert '<datalist id="registration-executive-authority-options">' in content
    assert '<datalist id="registration-institution-options">' in content
    assert '<option value="Департамент финансов города Москвы">ДФМ</option>' in content
    assert '<option value="ГБУ «Жилищник»">Жилищник</option>' in content


@pytest.mark.django_db
def test_form_page_renders_without_datalist_when_reference_empty(client):
    response = client.get(reverse('registration_request'))
    assert response.status_code == 200
    content = response.content.decode('utf-8')
    assert '<datalist' not in content


def test_fetch_legal_persons_dedupes_and_strips(monkeypatch):
    rows = [
        ('  ДФМ  ', ' Департамент финансов города Москвы '),
        ('ДФМ', 'Другая запись с тем же кратким именем'),
        ('', 'Пустое краткое имя'),
        (None, 'Нет краткого имени'),
        ('Пустое полное', '   '),
    ]

    class FakeCursor:
        def execute(self, sql, params=None):
            pass

        def fetchall(self):
            return rows

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

    class FakeConnection:
        def cursor(self):
            return FakeCursor()

    monkeypatch.setattr(registration_reference, 'connections', {'qgis': FakeConnection()})
    items = registration_reference._fetch_legal_persons('select 1')
    assert items == [{'shortname': 'ДФМ', 'fullname': 'Департамент финансов города Москвы'}]


def test_reference_cache_reuses_result_and_backs_off_on_failure(monkeypatch):
    registration_reference._cache.clear()
    calls = {'fetch': 0}

    def fake_fetch(sql):
        calls['fetch'] += 1
        return [{'shortname': 'A', 'fullname': 'AAA'}]

    monkeypatch.setattr(registration_reference, '_fetch_legal_persons', fake_fetch)
    assert registration_reference.list_executive_authorities() == [
        {'shortname': 'A', 'fullname': 'AAA'}
    ]
    assert registration_reference.list_executive_authorities() == [
        {'shortname': 'A', 'fullname': 'AAA'}
    ]
    assert calls['fetch'] == 1

    def failing_fetch(sql):
        calls['fetch'] += 1
        raise RuntimeError('db down')

    monkeypatch.setattr(registration_reference, '_fetch_legal_persons', failing_fetch)
    registration_reference._cache.clear()
    assert registration_reference.list_executive_authorities() == []
    assert registration_reference.list_executive_authorities() == []
    assert calls['fetch'] == 2
    registration_reference._cache.clear()
