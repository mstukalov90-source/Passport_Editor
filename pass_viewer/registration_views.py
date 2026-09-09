"""Публичная форма заявки на регистрацию и страницы просмотра заявок для МГГТ."""

import logging
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_POST

from .forms import RegistrationRequestForm
from .models import RegistrationRequest
from .roles import ROLE_MGGT, resolve_user_scope

logger = logging.getLogger(__name__)

# Заголовки и ширины колонок листа «Перечень» шаблона
# «Шаблон для добавления пользователей.xlsx».
XLSX_HEADERS = (
    "№ п/п",
    "Орган исполнительной власти",
    "Наименование учреждения",
    "ФИО ответственного представителя",
    "Должность",
    "Контактный телефон",
    "Адрес электронной почты",
)
XLSX_COLUMN_WIDTHS = (9.14, 28.43, 35.71, 36.14, 66.29, 37.86, 40.43)

_VALID_STATUSES = {RegistrationRequest.STATUS_NEW, RegistrationRequest.STATUS_PROCESSED}


def registration_request(request):
    """Публичная форма подачи заявки на регистрацию пользователя."""
    if request.method == "POST":
        form = RegistrationRequestForm(request.POST)
        if form.is_valid():
            RegistrationRequest.objects.create(**form.cleaned_data)
            logger.info("Новая заявка на регистрацию: %s", form.cleaned_data.get("email"))
            return redirect("registration_request_sent")
    else:
        form = RegistrationRequestForm()
    return render(request, "registration/registration_request.html", {"form": form})


def registration_request_sent(request):
    """Подтверждение отправки заявки."""
    return render(request, "registration/registration_request_sent.html")


def _mggt_scope(request):
    scope = resolve_user_scope(request.user.username)
    return scope if scope.role == ROLE_MGGT else None


def _filtered_requests(request):
    status = request.GET.get("status", "")
    queryset = RegistrationRequest.objects.all()
    if status in _VALID_STATUSES:
        queryset = queryset.filter(status=status)
    return queryset, status


def _list_url(filter_status):
    url = reverse("registration_requests_list")
    if filter_status in _VALID_STATUSES:
        url = f"{url}?status={filter_status}"
    return url


@login_required
@require_GET
def registration_requests_list(request):
    """Список заявок на регистрацию для сотрудников МГГТ."""
    scope = _mggt_scope(request)
    if scope is None:
        return HttpResponseForbidden("Доступ разрешён только сотрудникам МГГТ.")
    queryset, status = _filtered_requests(request)
    return render(
        request,
        "pass_viewer/registration_requests.html",
        {
            "requests": queryset,
            "status_filter": status,
            "new_count": RegistrationRequest.objects.filter(
                status=RegistrationRequest.STATUS_NEW
            ).count(),
            "user_role": scope.role,
        },
    )


@login_required
@require_POST
def registration_request_set_status(request, pk):
    """Пометить заявку обработанной или вернуть в новые."""
    if _mggt_scope(request) is None:
        return HttpResponseForbidden("Доступ разрешён только сотрудникам МГГТ.")
    new_status = request.POST.get("status", "")
    if new_status not in _VALID_STATUSES:
        return HttpResponseForbidden("Недопустимый статус заявки.")
    item = get_object_or_404(RegistrationRequest, pk=pk)
    item.status = new_status
    item.save(update_fields=["status"])
    return redirect(_list_url(request.POST.get("filter_status", "")))


@login_required
@require_GET
def registration_requests_export(request):
    """Выгрузка заявок в .xlsx по шаблону «Шаблон для добавления пользователей.xlsx»."""
    if _mggt_scope(request) is None:
        return HttpResponseForbidden("Доступ разрешён только сотрудникам МГГТ.")
    queryset, _ = _filtered_requests(request)

    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Перечень"
    sheet.append(list(XLSX_HEADERS))

    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in sheet[1]:
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border
    sheet.row_dimensions[1].height = 40

    for index, width in enumerate(XLSX_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for number, item in enumerate(queryset, start=1):
        text_values = [
            item.executive_authority,
            item.institution_name,
            item.representative_name,
            item.position,
            item.phone,
            item.email,
        ]
        sheet.append([number] + [ILLEGAL_CHARACTERS_RE.sub("", str(value)) for value in text_values])
        for cell in sheet[sheet.max_row]:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    buffer = BytesIO()
    workbook.save(buffer)
    filename = f"registration-requests-{timezone.now().date().isoformat()}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
